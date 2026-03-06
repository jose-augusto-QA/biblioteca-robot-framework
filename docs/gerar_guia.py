from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─── Dados das libraries ───────────────────────────────────────────────────────
libraries = [
    {
        "titulo": "1. robotframework 7.4.1 — Framework Principal",
        "descricao": (
            "Núcleo do Robot Framework. Responsável por fornecer a estrutura base para a criação, "
            "organização e execução de testes automatizados. Todas as demais bibliotecas dependem desta instalação."
        ),
        "topicos": [
            "Estrutura base para escrita de testes automatizados",
            "Suporte a keywords customizadas e bibliotecas externas",
            "Geração de relatórios e logs de execução",
            "Compatível com abordagens BDD, ATDD e keyword-driven",
        ],
    },
    {
        "titulo": "2. robotframework-assertion-engine 3.0.3 — Motor de Asserções",
        "descricao": (
            "Aprimora as mensagens de falha nas asserções, fornecendo diagnósticos detalhados sobre "
            "divergências entre valores esperados e obtidos, facilitando a análise de falhas."
        ),
        "topicos": [
            "Mensagens de erro detalhadas e legíveis",
            "Suporte a comparações complexas de valores",
            "Integração nativa com outras libraries do ecossistema",
        ],
    },
    {
        "titulo": "3. robotframework-databaselibrary 2.4.1 — Integração com Banco de Dados",
        "descricao": (
            "Possibilita a conexão e execução de comandos SQL em bancos de dados relacionais, "
            "incluindo MySQL, PostgreSQL, SQLite e Oracle. Indicada para validações de persistência de dados."
        ),
        "topicos": [
            "Conexão com múltiplos bancos de dados relacionais",
            "Execução de queries SQL (SELECT, INSERT, UPDATE, DELETE)",
            "Validação de dados persistidos após ações do sistema",
            "Suporte a transações e rollback",
        ],
    },
    {
        "titulo": "4. robotframework-datadriver 1.11.2 — Testes Orientados a Dados",
        "descricao": (
            "Permite a parametrização de testes a partir de fontes de dados externas, como arquivos "
            ".csv, .xlsx e .json, viabilizando a execução do mesmo caso de teste com múltiplos conjuntos de dados."
        ),
        "topicos": [
            "Leitura de dados de arquivos CSV, Excel e JSON",
            "Execução automática do mesmo teste com diferentes entradas",
            "Redução de duplicidade nos casos de teste",
            "Ideal para testes de validação de formulários e regras de negócio",
        ],
    },
    {
        "titulo": "5. robotframework-excellib 2.0.1 — Manipulação de Planilhas Excel",
        "descricao": (
            "Oferece recursos para leitura e escrita de arquivos .xlsx. Diferencia-se do DataDriver "
            "por permitir também a gravação de resultados em planilhas, sendo útil em relatórios customizados."
        ),
        "topicos": [
            "Leitura e escrita de arquivos .xlsx",
            "Criação de relatórios customizados em Excel",
            "Manipulação de células, linhas e colunas",
            "Utilizado em conjunto com DataDriver para cenários avançados",
        ],
    },
    {
        "titulo": "6. robotframework-faker 6.0.0 — Geração de Dados Sintéticos",
        "descricao": (
            "Gera dados fictícios e realistas para uso em testes, tais como nomes, endereços, e-mails, "
            "CPFs e datas, eliminando a necessidade de massa de dados manual."
        ),
        "topicos": [
            "Geração de nomes, e-mails, CPFs e endereços fictícios",
            "Suporte a localização (pt_BR) para dados regionais",
            "Elimina dependência de massa de dados estática",
            "Integração direta com variáveis do Robot Framework",
        ],
    },
    {
        "titulo": "7. robotframework-jsonlibrary 0.5 — Manipulação de JSON",
        "descricao": (
            "Fornece recursos para leitura, validação e manipulação de estruturas JSON. "
            "Amplamente utilizada em testes de APIs REST para verificação de payloads de resposta."
        ),
        "topicos": [
            "Leitura e parsing de arquivos e strings JSON",
            "Validação de estrutura e valores de respostas de API",
            "Suporte a JSONPath para navegação em objetos complexos",
            "Conversão entre JSON e dicionários Python",
        ],
    },
    {
        "titulo": "8. robotframework-lsp 1.13.0 — Suporte ao Ambiente de Desenvolvimento",
        "descricao": (
            "Implementa o protocolo LSP (Language Server Protocol), habilitando recursos de produtividade "
            "no editor de código, como autocomplete, realce de sintaxe e navegação entre keywords. "
            "Utilizado exclusivamente no ambiente de desenvolvimento."
        ),
        "topicos": [
            "Autocomplete de keywords e variáveis no VSCode",
            "Realce de sintaxe para arquivos .robot e .resource",
            "Navegação entre definições de keywords",
            "Não impacta a execução dos testes",
        ],
    },
    {
        "titulo": "9. robotframework-pabot 5.2.2 — Execução Paralela",
        "descricao": (
            "Viabiliza a execução simultânea de múltiplos testes, reduzindo significativamente o tempo "
            "total de execução das suítes. Recomendado para pipelines de integração contínua."
        ),
        "topicos": [
            "Execução paralela de suítes e casos de teste",
            "Redução expressiva do tempo de execução",
            "Suporte a compartilhamento de recursos entre processos",
            "Integração com pipelines CI/CD (GitHub Actions, Jenkins)",
        ],
    },
    {
        "titulo": "10. robotframework-pythonlibcore 4.5.0 — Núcleo para Libraries Python",
        "descricao": (
            "Biblioteca base utilizada como dependência por outras libraries, como a SeleniumLibrary. "
            "Fornece a estrutura necessária para o desenvolvimento de bibliotecas customizadas em Python."
        ),
        "topicos": [
            "Dependência técnica de outras libraries do ecossistema",
            "Estrutura base para criação de libraries customizadas",
            "Geralmente instalado automaticamente como dependência",
        ],
    },
    {
        "titulo": "11. robotframework-requests 0.9.7 — Testes de APIs REST",
        "descricao": (
            "Permite a execução de requisições HTTP (GET, POST, PUT, DELETE, PATCH) para validação de "
            "APIs REST. Principal biblioteca para automação de testes de serviços web."
        ),
        "topicos": [
            "Suporte a todos os métodos HTTP (GET, POST, PUT, DELETE, PATCH)",
            "Envio de headers, autenticação e body nas requisições",
            "Validação de status code e payload de resposta",
            "Integração com JsonLibrary para validação de respostas JSON",
        ],
    },
    {
        "titulo": "12. robotframework-seleniumlibrary 6.8.0 — Automação Web",
        "descricao": (
            "Possibilita a automação de interações em navegadores web, incluindo preenchimento de "
            "formulários, cliques, navegação entre páginas e captura de screenshots. "
            "Principal biblioteca para testes de interface web."
        ),
        "topicos": [
            "Suporte a múltiplos navegadores (Chrome, Firefox, Edge)",
            "Interação com elementos web (cliques, inputs, selects)",
            "Captura de screenshots em caso de falha",
            "Suporte a iframes, janelas múltiplas e alerts",
            "Integração com Selenium WebDriver",
        ],
    },
]

# ─── Funções auxiliares ────────────────────────────────────────────────────────
def add_horizontal_line(paragraph):
    """Adiciona uma linha horizontal abaixo do parágrafo."""
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_cell_background(cell, color_hex):
    """Define cor de fundo de uma célula de tabela."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


# ─── Criação do documento ──────────────────────────────────────────────────────
doc = Document()

# Margens
section = doc.sections[0]
section.top_margin    = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin   = Inches(1.2)
section.right_margin  = Inches(1.2)

# ── Título principal ──
titulo_principal = doc.add_heading("Guia de Libraries do Robot Framework", level=0)
titulo_principal.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_titulo = titulo_principal.runs[0]
run_titulo.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
run_titulo.font.size = Pt(22)

doc.add_paragraph()

# ── Introdução ──
intro_heading = doc.add_heading("Introdução", level=1)
intro_heading.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

intro_text = doc.add_paragraph(
    "Este documento apresenta as principais bibliotecas utilizadas no ecossistema Robot Framework, "
    "descrevendo suas respectivas finalidades, casos de uso e exemplos de aplicação. "
    "O objetivo é fornecer uma referência técnica clara e objetiva para equipes de qualidade de software."
)
intro_text.paragraph_format.space_after = Pt(12)

add_horizontal_line(intro_text)
doc.add_paragraph()

# ── Tabela resumo ──
resumo_heading = doc.add_heading("Resumo das Libraries", level=1)
resumo_heading.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

tabela = doc.add_table(rows=1, cols=3)
tabela.style = "Table Grid"

# Cabeçalho
hdr_cells = tabela.rows[0].cells
headers = ["Library", "Versão", "Finalidade"]
cores_header = ["1F497D", "1F497D", "1F497D"]

for i, (cell, header) in enumerate(zip(hdr_cells, headers)):
    set_cell_background(cell, cores_header[i])
    run = cell.paragraphs[0].add_run(header)
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.size = Pt(10)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Dados da tabela
resumo_data = [
    ("robotframework",                "7.4.1",  "Framework principal"),
    ("robotframework-assertion-engine","3.0.3", "Motor de asserções aprimorado"),
    ("robotframework-databaselibrary","2.4.1",  "Integração com banco de dados"),
    ("robotframework-datadriver",     "1.11.2", "Testes orientados a dados"),
    ("robotframework-excellib",       "2.0.1",  "Manipulação de planilhas Excel"),
    ("robotframework-faker",          "6.0.0",  "Geração de dados sintéticos"),
    ("robotframework-jsonlibrary",    "0.5",    "Manipulação de JSON"),
    ("robotframework-lsp",            "1.13.0", "Suporte ao ambiente de desenvolvimento"),
    ("robotframework-pabot",          "5.2.2",  "Execução paralela de testes"),
    ("robotframework-pythonlibcore",  "4.5.0",  "Núcleo para libraries Python"),
    ("robotframework-requests",       "0.9.7",  "Testes de APIs REST"),
    ("robotframework-seleniumlibrary","6.8.0",  "Automação de testes Web"),
]

for i, (lib, versao, finalidade) in enumerate(resumo_data):
    row_cells = tabela.add_row().cells
    bg_color = "EBF3FB" if i % 2 == 0 else "FFFFFF"
    set_cell_background(row_cells[0], bg_color)
    set_cell_background(row_cells[1], bg_color)
    set_cell_background(row_cells[2], bg_color)

    r0 = row_cells[0].paragraphs[0].add_run(lib)
    r0.font.size = Pt(9)
    r0.font.name = "Courier New"

    r1 = row_cells[1].paragraphs[0].add_run(versao)
    r1.font.size = Pt(9)
    row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    r2 = row_cells[2].paragraphs[0].add_run(finalidade)
    r2.font.size = Pt(9)

doc.add_paragraph()

# ── Detalhamento de cada library ──
detalhes_heading = doc.add_heading("Detalhamento das Libraries", level=1)
detalhes_heading.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

for lib in libraries:
    # Título da library
    h2 = doc.add_heading(lib["titulo"], level=2)
    h2.runs[0].font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    h2.runs[0].font.size = Pt(12)

    # Descrição
    desc_label = doc.add_paragraph()
    run_label = desc_label.add_run("Descrição: ")
    run_label.bold = True
    run_label.font.size = Pt(10)
    run_desc = desc_label.add_run(lib["descricao"])
    run_desc.font.size = Pt(10)
    desc_label.paragraph_format.space_after = Pt(4)

    # Tópicos
    topicos_label = doc.add_paragraph()
    run_top_label = topicos_label.add_run("Características:")
    run_top_label.bold = True
    run_top_label.font.size = Pt(10)
    topicos_label.paragraph_format.space_after = Pt(2)

    for topico in lib["topicos"]:
        bullet = doc.add_paragraph(style="List Bullet")
        run_bullet = bullet.add_run(topico)
        run_bullet.font.size = Pt(10)
        bullet.paragraph_format.space_after = Pt(2)
        bullet.paragraph_format.left_indent = Inches(0.3)

    # Linha separadora
    sep = doc.add_paragraph()
    add_horizontal_line(sep)
    doc.add_paragraph()

# ── Rodapé ──
section = doc.sections[0]
footer = section.footer
footer_para = footer.paragraphs[0]
footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_footer = footer_para.add_run("Guia de Libraries — Robot Framework  |  Biblioteca de Apoio à Automação")
run_footer.font.size = Pt(8)
run_footer.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

# ── Salvar ──
nome_arquivo = "Guia_Libraries_RobotFramework.docx"
doc.save(nome_arquivo)
print(f'✅ Arquivo "{nome_arquivo}" gerado com sucesso!')
